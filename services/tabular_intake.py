"""Import invoice rows from CSV/XLSX without pretending they came from OCR."""
from __future__ import annotations

import csv
import io
import zipfile
from datetime import date, datetime, timezone
from pathlib import PurePosixPath

from openpyxl import load_workbook

from models.invoice_schema import InvoiceResult
from services.validation_service import validate_invoice

MAX_ROWS = 10_000
MAX_COLUMNS = 100
MAX_EXPANDED_BYTES = 100_000_000

FIELD_ALIASES = {
    'invoice_number': ('invoice_number', 'Invoice Number (As per Invoice)', 'Invoice Number', 'Invoice'),
    'invoice_date': ('invoice_date', 'Invoice Date', 'Date'),
    'vendor_name': ('vendor_name', 'Vendor -As per MDM', 'Vendor'),
    'customer_name': ('customer_name', 'Customer'),
    'currency': ('currency', 'Currency'),
    'subtotal': ('subtotal', 'Net Amount', 'Subtotal'),
    'tax_amount': ('tax_amount', 'Tax Amount'),
    'total_amount': ('total_amount', 'Invoice Amount', 'Total Amount', 'Total'),
    'service_start_date': ('service_start_date', 'Billing Start Date', 'Service Start Date'),
    'service_end_date': ('service_end_date', 'Billing End Date', 'Service End Date'),
    'billing_period': ('billing_period', 'Period'),
    'bank_account_number': ('bank_account_number', 'Billing Account', 'MDM Account number'),
    'customer_email': ('customer_email', 'TASK - Email'),
    'vendor_email': ('vendor_email', 'Vendor Email', 'Source email'),
    'due_date': ('due_date', 'Due Date'),
    'region': ('region', 'Region'),
    'scope': ('scope', 'Scope'),
    'invoice_comments': ('invoice_comments', 'Comments'),
}


def _value(value: object) -> object:
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return value


def _read_csv(data: bytes):
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError as exc:
        raise ValueError('CSV must be UTF-8 encoded') from exc
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise ValueError('CSV has no header row')
    rows = []
    try:
        for number, row in enumerate(reader, 2):
            if number > MAX_ROWS + 1:
                raise ValueError(f'CSV exceeds {MAX_ROWS} data rows')
            rows.append(row)
    except csv.Error as exc:
        raise ValueError('CSV contains an invalid row') from exc
    return reader.fieldnames, rows


def _read_xlsx(data: bytes):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if sum(item.file_size for item in members) > MAX_EXPANDED_BYTES or any(
                item.file_size > MAX_EXPANDED_BYTES or
                (item.compress_size and item.file_size / item.compress_size > 100)
                for item in members
            ):
                raise ValueError('Excel workbook expands beyond the safety limit')
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Excel workbook is invalid or unreadable') from exc
    try:
        sheet = book['Invoice_Data'] if 'Invoice_Data' in book.sheetnames else book.active
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if not header:
            raise ValueError('Excel worksheet has no header row')
        if len(header) > MAX_COLUMNS:
            raise ValueError('Excel worksheet has too many columns')
        names = [str(value).strip() if value is not None else '' for value in header]
        rows = []
        for number, values in enumerate(iterator, 2):
            if number > MAX_ROWS + 1:
                raise ValueError(f'Excel worksheet exceeds {MAX_ROWS} data rows')
            rows.append(dict(zip(names, values)))
        return names, rows
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Excel worksheet is invalid or unreadable') from exc
    finally:
        book.close()


def import_invoice_rows(filename: str, data: bytes, tolerance: float, threshold: float) -> list[InvoiceResult]:
    """Map known invoice columns into independently reviewable saved records."""
    safe_name = PurePosixPath(filename.replace('\\', '/')).name
    suffix = PurePosixPath(safe_name).suffix.lower()
    if suffix == '.csv':
        names, rows = _read_csv(data)
    elif suffix == '.xlsx':
        names, rows = _read_xlsx(data)
    else:
        raise ValueError('Only UTF-8 CSV and XLSX spreadsheets are supported')
    if len(names) > MAX_COLUMNS or len(rows) > MAX_ROWS:
        raise ValueError(f'Spreadsheet exceeds the {MAX_ROWS}-row or {MAX_COLUMNS}-column limit')
    if len(set(names)) != len(names) or any(not name for name in names):
        raise ValueError('Spreadsheet has empty or repeated column names')
    available = set(names)
    if not any(alias in available for alias in FIELD_ALIASES['invoice_number']):
        raise ValueError('Spreadsheet needs an Invoice or Invoice Number column')
    if not any(alias in available for alias in FIELD_ALIASES['vendor_name']):
        raise ValueError('Spreadsheet needs a Vendor column')
    results = []
    for row_number, row in enumerate(rows, 2):
        if None in row:
            raise ValueError(f'Row {row_number} has more values than header columns')
        if not any(value not in (None, '') for value in row.values()):
            continue
        values = {}
        for field, aliases in FIELD_ALIASES.items():
            for alias in aliases:
                value = _value(row.get(alias))
                if value not in (None, ''):
                    values[field] = value
                    break
        source = f'{safe_name}#row-{row_number}'
        result = InvoiceResult.from_ai_payload(
            {'invoice': values, 'line_items': [], 'field_confidence': {}},
            source, datetime.now(timezone.utc).isoformat(), threshold)
        result.processing_status = 'IMPORTED'
        validate_invoice(result, tolerance, threshold)
        result.review_required = True
        result.validation_status = 'REVIEW_REQUIRED'
        result.warnings.append('Imported spreadsheet row; verify values against the original invoice.')
        results.append(result)
    if not results:
        raise ValueError('Spreadsheet has no nonempty invoice rows')
    return results
