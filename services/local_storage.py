"""Durable batch history with recoverable, cumulative CSV/XLSX exports.

SQLite is authoritative; exports are regenerated, never parsed back as history.
Each export is replaced atomically. A locked Excel file cannot lose saved data.
"""
from __future__ import annotations

import csv
import io
import json
import math
import os
import re
import sqlite3
import tempfile
import threading
from functools import wraps
from uuid import uuid4
from contextlib import contextmanager
from dataclasses import asdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill

from models.invoice_schema import Invoice, InvoiceResult, LineItem
from services.csv_service import generate_csvs
from services.storage_lock import exclusive_folder_operation
from services.mdm_csv_service import REPORT_NAME as MDM_REPORT_NAME, generate_mdm_csv


_STORAGE_LOCK = threading.RLock()


def _current_storage(method):
    @wraps(method)
    def guarded(self, *args, **kwargs):
        with _STORAGE_LOCK:
            if not self.database.is_file():
                raise RuntimeError('Local data was reset. Reload the page before saving.')
            with self._connection() as connection:
                identity = connection.execute('SELECT token FROM storage_identity WHERE id=1').fetchone()
            if not identity or identity[0] != self.generation:
                raise RuntimeError('Local data was reset. Reload the page before saving.')
            return method(self, *args, **kwargs)
    return guarded


def _atomic_write(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix='.pending-', dir=path.parent)
    try:
        with os.fdopen(descriptor, 'wb') as stream:
            stream.write(payload)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def _safe_filename(name: str) -> str:
    name = name.replace('\\', '/').rsplit('/', 1)[-1]
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', name)[:140].strip(' .') or 'upload'


def _restore(raw: dict) -> InvoiceResult:
    raw = dict(raw)
    raw['invoice'] = Invoice.from_dict(raw['invoice'])
    raw['line_items'] = [LineItem.from_dict(item) for item in raw['line_items']]
    return InvoiceResult(**raw)


def _excel_bytes(reports: dict[str, bytes]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for filename, payload in reports.items():
        sheet = workbook.create_sheet(
            'MDM Invoice Data' if filename == MDM_REPORT_NAME else Path(filename).stem)
        rows = list(csv.reader(io.StringIO(payload.decode('utf-8-sig'))))
        if len(rows) > 1_048_576:
            raise ValueError('Excel row limit exceeded. CSV and database history remain available.')
        numeric_columns = {'Invoice Amount', 'Net Amount', 'Tax Amount', 'Ageing (In Days)',
                           'Service line count', 'Total quantity',
                           'Extraction completeness (%)',
                           'Variance %', 'Variance amounttTax', 'Invoice Aging',
                           'quantity', 'unit_price', 'tax', 'discount', 'line_total', 'confidence',
                           'line_number', 'attachment_index', 'record_index'}
        for row_index, row in enumerate(rows):
            # Text cells preserve account numbers/leading zeros, and prevent
            # invoice text beginning with = from becoming executable formulas.
            sheet.append([ILLEGAL_CHARACTERS_RE.sub('', value) for value in row])
            for cell in sheet[sheet.max_row]:
                cell.data_type = 's'
                if row_index and rows[0][cell.column - 1] in numeric_columns and cell.value:
                    try:
                        number = float(cell.value)
                        if math.isfinite(number):
                            cell.value = number
                            cell.number_format = '0.00' if rows[0][cell.column - 1] in {
                                'Invoice Amount', 'Net Amount', 'Tax Amount',
                                'Variance %', 'Variance amounttTax',
                                'unit_price', 'tax', 'discount', 'line_total', 'confidence',
                                'Extraction completeness (%)'} else '0.########'
                    except ValueError:
                        pass
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='234E70')
            sheet.column_dimensions[cell.column_letter].width = min(40, max(18, len(str(cell.value)) + 2))
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class LocalStorage:
    def __init__(self, root: Path | str):
        self.root = Path(root).expanduser().resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        self.database = self.root / 'invoice_history.sqlite3'
        with _STORAGE_LOCK, self._connection() as connection:
            connection.execute('CREATE TABLE IF NOT EXISTS batches ('
                               'batch_id TEXT PRIMARY KEY, results TEXT NOT NULL, '
                               'created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)')
            connection.execute('CREATE TABLE IF NOT EXISTS export_state (id INTEGER PRIMARY KEY, warnings TEXT NOT NULL)')
            connection.execute('CREATE TABLE IF NOT EXISTS storage_identity (id INTEGER PRIMARY KEY, token TEXT NOT NULL)')
            connection.execute('CREATE TABLE IF NOT EXISTS outlook_intake ('
                               'filename TEXT NOT NULL, sha256 TEXT NOT NULL, batch_id TEXT NOT NULL, '
                               'processed_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, '
                               'PRIMARY KEY(filename, sha256))')
            connection.execute('CREATE TABLE IF NOT EXISTS reviewed_examples ('
                               'batch_id TEXT NOT NULL, result_index INTEGER NOT NULL, '
                               'ocr_text TEXT NOT NULL, invoice_json TEXT NOT NULL, '
                               'reviewed_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP, '
                               'PRIMARY KEY(batch_id, result_index))')
            connection.execute('INSERT OR IGNORE INTO storage_identity VALUES (1, ?)', (uuid4().hex,))
            self.generation = connection.execute('SELECT token FROM storage_identity WHERE id=1').fetchone()[0]

    @contextmanager
    def _connection(self):
        connection = sqlite3.connect(self.database, timeout=30)
        try:
            with connection:
                yield connection
        finally:
            connection.close()

    def _batch_path(self, batch_id: str) -> Path:
        if not re.fullmatch(r'[0-9a-f]{32}', batch_id):
            raise ValueError('Invalid batch identifier')
        return self.root / 'uploads' / batch_id

    @_current_storage
    def archive_uploads(self, batch_id: str, uploads: list[tuple[str, bytes]]) -> None:
        directory = self._batch_path(batch_id) / 'originals'
        manifest = []
        for index, (name, data) in enumerate(uploads, 1):
            stored_name = f'{index:04d}_{_safe_filename(name)}'
            _atomic_write(directory / stored_name, data)
            manifest.append({'original_filename': name, 'stored_filename': stored_name})
        _atomic_write(directory / 'manifest.json', json.dumps(manifest, ensure_ascii=False, indent=2).encode())

    @_current_storage
    def archive_documents(self, batch_id: str, documents: list) -> None:
        directory = self._batch_path(batch_id) / 'documents'
        manifest = []
        for index, document in enumerate(documents, 1):
            if document.intake_status or not document.data:
                continue
            stored_name = f'{index:04d}_{_safe_filename(document.filename)}'
            _atomic_write(directory / stored_name, document.data)
            manifest.append({'original_filename': document.filename, 'stored_filename': stored_name,
                             'source_email': document.source_email,
                             'source_attachment': document.source_attachment,
                             'attachment_index': document.attachment_index})
        _atomic_write(directory / 'manifest.json', json.dumps(manifest, ensure_ascii=False, indent=2).encode())

    @_current_storage
    def save_batch(self, batch_id: str, results: list[InvoiceResult],
                   outlook_files: list[tuple[str, str]] | None = None) -> list[str]:
        self._batch_path(batch_id)
        payload = json.dumps([asdict(result) for result in results], default=str, ensure_ascii=False)
        with self._connection() as connection:
            connection.execute('INSERT INTO batches(batch_id, results) VALUES (?, ?) '
                               'ON CONFLICT(batch_id) DO UPDATE SET results=excluded.results', (batch_id, payload))
            connection.execute('INSERT OR REPLACE INTO export_state VALUES (1, ?)',
                               (json.dumps(['Reports need refreshing from the saved database.']),))
            for filename, digest in outlook_files or []:
                connection.execute('INSERT OR IGNORE INTO outlook_intake(filename, sha256, batch_id) '
                                   'VALUES (?, ?, ?)', (filename, digest, batch_id))
        return self.refresh_exports()

    @property
    @_current_storage
    def outlook_folder(self) -> Path:
        folder = self.root / 'Outlook'
        if folder.is_symlink() or folder.is_junction():
            raise ValueError('Outlook intake folder cannot be a link.')
        folder.mkdir(exist_ok=True)
        return folder

    @_current_storage
    def outlook_processed(self, filename: str, digest: str) -> bool:
        with self._connection() as connection:
            return connection.execute('SELECT 1 FROM outlook_intake WHERE filename=? AND sha256=?',
                                      (filename, digest)).fetchone() is not None

    @_current_storage
    def export_warnings(self) -> list[str]:
        with self._connection() as connection:
            row = connection.execute('SELECT warnings FROM export_state WHERE id=1').fetchone()
            return json.loads(row[0]) if row else []

    @_current_storage
    def load_batches(self) -> list[tuple[str, list[InvoiceResult]]]:
        with self._connection() as connection:
            return [(batch_id, [_restore(item) for item in json.loads(payload)]) for batch_id, payload in
                    connection.execute('SELECT batch_id, results FROM batches ORDER BY rowid')]

    def count(self) -> tuple[int, int]:
        batches = self.load_batches()
        return len(batches), sum(len(results) for _, results in batches)

    @_current_storage
    def training_examples(self) -> list[tuple[str, dict]]:
        with self._connection() as connection:
            return [(text, json.loads(invoice)) for text, invoice in connection.execute(
                'SELECT ocr_text, invoice_json FROM reviewed_examples ORDER BY reviewed_at')]

    @_current_storage
    def remember_review(self, batch_id: str, result_index: int, result: InvoiceResult) -> bool:
        """Save only human-reviewed attachment data; never email bodies or unreviewed guesses."""
        from services.ocr_service import extract_document_text
        directory = self._batch_path(batch_id) / 'documents'
        manifest_path = directory / 'manifest.json'
        if not manifest_path.is_file():
            return False
        manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
        candidates = [item for item in manifest
                      if item['original_filename'] == result.source_file
                      and item.get('source_email', '') == result.source_email
                      and item.get('attachment_index') == result.attachment_index]
        if len(candidates) != 1:
            return False
        stored = directory / candidates[0]['stored_filename']
        if not stored.resolve().is_relative_to(directory.resolve()):
            return False
        try:
            ocr_text = extract_document_text(result.source_file, stored.read_bytes())
        except (OSError, ValueError, RuntimeError):
            return False
        with self._connection() as connection:
            connection.execute('INSERT INTO reviewed_examples(batch_id, result_index, ocr_text, invoice_json) '
                               'VALUES (?, ?, ?, ?) ON CONFLICT(batch_id, result_index) '
                               'DO UPDATE SET ocr_text=excluded.ocr_text, invoice_json=excluded.invoice_json, '
                               'reviewed_at=CURRENT_TIMESTAMP',
                               (batch_id, result_index, ocr_text[:30000],
                                json.dumps(result.invoice.to_dict(), ensure_ascii=False)))
        return True

    @_current_storage
    def refresh_exports(self) -> list[str]:
        warnings = []
        # Serialize exporters across sessions/processes and take a consistent
        # snapshot. Data was committed before this export-only transaction.
        with self._connection() as connection:
            connection.execute('BEGIN IMMEDIATE')
            combined = {}
            headers = {}
            all_results = []
            for batch_id, payload in connection.execute('SELECT batch_id, results FROM batches ORDER BY rowid'):
                results = [_restore(item) for item in json.loads(payload)]
                all_results.extend(results)
                for index, result in enumerate(results, 1):
                    for filename, data in generate_csvs([result]).items():
                        reader = csv.DictReader(io.StringIO(data.decode()))
                        headers[filename] = list(reader.fieldnames) + ['batch_id', 'record_index']
                        if 'source_file' not in headers[filename]:
                            headers[filename].append('source_file')
                        for row in reader:
                            row.update(batch_id=batch_id, record_index=index, source_file=result.source_file)
                            combined.setdefault(filename, []).append(row)
            reports = {}
            for filename, empty in generate_csvs([]).items():
                default = next(csv.reader(io.StringIO(empty.decode()))) + ['batch_id', 'record_index']
                if 'source_file' not in default:
                    default.append('source_file')
                output = io.StringIO(newline='')
                writer = csv.DictWriter(output, fieldnames=headers.get(filename, default))
                writer.writeheader()
                writer.writerows(combined.get(filename, []))
                reports[filename] = output.getvalue().encode('utf-8-sig')
            reports[MDM_REPORT_NAME] = generate_mdm_csv(all_results)
            for filename, payload in reports.items():
                try:
                    _atomic_write(self.root / filename, payload)
                except OSError as exc:
                    warnings.append(f'{filename}: {exc}')
            try:
                _atomic_write(self.root / 'Invoice_History.xlsx', _excel_bytes(reports))
            except (OSError, ValueError) as exc:
                warnings.append(f'Invoice_History.xlsx: {exc}')
            connection.execute('INSERT OR REPLACE INTO export_state VALUES (1, ?)', (json.dumps(warnings),))
        return warnings

    @_current_storage
    def report_files(self) -> dict[str, bytes]:
        names = ['Invoice_History.xlsx', 'Invoice_Data.csv', 'Invoice_Line_Items.csv', 'Audit_Report.csv', MDM_REPORT_NAME]
        return {name: (self.root / name).read_bytes() for name in names if (self.root / name).is_file()}

    @_current_storage
    @exclusive_folder_operation
    def reset(self, confirmed_root: str) -> dict:
        """Clear invoice history and uploads while retaining learned training data.

        The SQLite database, reviewed examples, and models stay in place. A new
        storage identity invalidates tabs opened before reset. No recursive
        delete follows a symlink or Windows junction.
        """
        root = self.root
        if str(root) != confirmed_root or root.resolve() != root or root.is_symlink() or root.is_junction():
            raise ValueError('Storage location changed or is a link. Reset refused.')
        project = Path(__file__).resolve().parents[1]
        protected = [Path.home().resolve(), project, Path(root.anchor)]
        protected += [Path(os.environ[key]).resolve() for key in ('SystemRoot', 'ProgramFiles', 'ProgramFiles(x86)') if os.environ.get(key)]
        if any(root == item or root in item.parents for item in protected):
            raise ValueError('Reset cannot target a home, project, system, or drive root.')
        files, directories = [], []
        intake_roots = (root / 'uploads', root / 'Outlook')
        def unreadable(error):
            raise error
        # Validate the entire tree before deleting anything, including nested
        # paths, so an unexpected link fails without deleting earlier files.
        for directory, names, filenames in os.walk(root, followlinks=False, onerror=unreadable):
            for name in names + filenames:
                target = Path(directory) / name
                if target.is_symlink() or target.is_junction() or not target.resolve().is_relative_to(root):
                    raise ValueError(f'Reset refused: linked or external path {target.name}')
                if any(target != intake and target.is_relative_to(intake) for intake in intake_roots):
                    (directories if target.is_dir() else files).append(target)
        count = 0
        try:
            # Keep the database intact until uploaded files are gone. If a
            # source is locked, its saved result remains available for retry.
            for target in files:
                target.unlink()
                count += 1
            for target in sorted(directories, key=lambda path: len(path.parts), reverse=True):
                target.rmdir()
        except OSError as exc:
            raise OSError(f'Reset incomplete: {count} files deleted. Close Excel/other app instances and retry. {exc}') from exc
        cleared_records = self.count()[1]
        with self._connection() as connection:
            connection.execute('DELETE FROM batches')
            connection.execute('DELETE FROM outlook_intake')
            connection.execute('DELETE FROM export_state')
            connection.execute('UPDATE storage_identity SET token=? WHERE id=1', (uuid4().hex,))
        fresh = LocalStorage(root)
        warnings = fresh.refresh_exports()
        return {'deleted_files': count, 'cleared_records': cleared_records,
                'retained_training_examples': len(fresh.training_examples()),
                'warnings': warnings}
