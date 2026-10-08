from dataclasses import asdict
from contextlib import closing
import io
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest

from openpyxl import load_workbook
from evaluation.reporting import export_run, status_counts, invoice_rows
from models.invoice_schema import InvoiceResult


class ReportingTests(unittest.TestCase):
    def test_disjoint_status_counts_and_review_comments(self):
        result = InvoiceResult(errors=['Line-item mismatch on line 1'], warnings=['Low-confidence extraction (90%)'])
        records = [{'id': 'a.jpg', 'split': 'development', 'seconds': 1, 'result': asdict(result), 'scores': []},
                   {'id': 'b.jpg', 'split': 'test', 'seconds': 1, 'result': asdict(InvoiceResult.failed('b.jpg', '503', '')), 'scores': []}]
        self.assertEqual(status_counts(records), {'VALID': 0, 'REVIEW_REQUIRED': 1, 'FAILED': 1, 'SKIPPED': 0})
        self.assertIn('Line-item mismatch', invoice_rows(records)[0]['review_comments'])

    def test_all_exports_and_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory) / 'run'
            folder.mkdir()
            with closing(sqlite3.connect(folder / 'predictions.sqlite3')) as connection:
                connection.execute('CREATE TABLE predictions(id TEXT, split TEXT, payload TEXT)')
                record = {'id': 'a.jpg', 'split': 'test', 'seconds': 1, 'result': asdict(InvoiceResult()), 'scores': []}
                connection.execute('INSERT INTO predictions VALUES (?, ?, ?)', ('a.jpg', 'test', json.dumps(record)))
                connection.commit()
            metadata, records, reports = export_run(folder)
            self.assertEqual(len(records), 1)
            self.assertIn('All_Results.json', reports)
            book = load_workbook(io.BytesIO(reports['Test_Results.xlsx']))
            self.assertIn('Review_Reasons', book.sheetnames)
            self.assertIn('Status_Summary', book.sheetnames)
            book.close()
