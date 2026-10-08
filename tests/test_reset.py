"""Reset tests never touch the configured real local_data folder."""
import csv
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from uuid import uuid4

from openpyxl import load_workbook

from services.local_storage import LocalStorage
from models.invoice_schema import Invoice, InvoiceResult


class ResetTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name) / 'local_data'
        self.store = LocalStorage(self.root)
        self.root = self.store.root
        batch = uuid4().hex
        self.store.archive_uploads(batch, [('test.pdf', b'test')])
        self.store.save_batch(batch, [InvoiceResult(source_file='test.pdf')])
        self.store.outlook_folder.joinpath('message.msg').write_bytes(b'mail')
        model = self.root / 'models' / 'invoice_line_reviews.joblib'
        model.parent.mkdir(exist_ok=True)
        model.write_bytes(b'trained-model')
        with self.store._connection() as connection:
            connection.execute('INSERT INTO reviewed_examples(batch_id, result_index, ocr_text, invoice_json) '
                               'VALUES (?, ?, ?, ?)', (batch, 0, 'reviewed OCR', '{}'))

    def test_reset_clears_invoice_rows_but_keeps_training_and_rejects_stale_writers(self):
        sibling = self.root.parent / 'keep.txt'
        sibling.write_text('keep')
        outcome = self.store.reset(str(self.root))
        self.assertGreaterEqual(outcome['deleted_files'], 2)
        self.assertEqual(outcome['cleared_records'], 1)
        self.assertEqual(outcome['retained_training_examples'], 1)
        self.assertEqual(outcome['warnings'], [])
        self.assertEqual(list((self.root / 'uploads').iterdir()), [])
        self.assertEqual(list((self.root / 'Outlook').iterdir()), [])
        self.assertEqual((self.root / 'models' / 'invoice_line_reviews.joblib').read_bytes(), b'trained-model')
        self.assertEqual(sibling.read_text(), 'keep')
        fresh = LocalStorage(self.root)
        self.assertEqual(fresh.count(), (0, 0))
        self.assertEqual(len(fresh.training_examples()), 1)
        self.assertEqual(len((self.root / 'MDM' / 'Invoice_data.csv').read_text(encoding='utf-8-sig').splitlines()), 1)
        book = load_workbook(self.root / 'Invoice_History.xlsx', read_only=True)
        self.assertEqual(book['MDM Invoice Data'].max_row, 1)
        book.close()
        with self.assertRaises(RuntimeError):
            self.store.save_batch(uuid4().hex, [InvoiceResult()])
        fresh.save_batch(uuid4().hex, [InvoiceResult()])
        self.assertEqual(fresh.count(), (1, 1))

    def test_incorrect_confirmation_cannot_delete(self):
        with self.assertRaises(ValueError):
            self.store.reset(str(self.root.parent))
        self.assertEqual(self.store.count(), (1, 1))

    def test_cumulative_invoice_data_rows_append_then_reset_to_header(self):
        report = self.root / 'MDM' / 'Invoice_data.csv'
        second = InvoiceResult(invoice=Invoice(invoice_number='INV-2', vendor_name='Test Vendor'),
                               source_file='second.pdf')
        self.store.save_batch(uuid4().hex, [second])
        with report.open(encoding='utf-8-sig') as stream:
            before = list(csv.DictReader(stream))
        self.assertEqual(len(before), 2)
        self.assertEqual(before[1]['Invoice Number'], 'INV-2')
        self.assertEqual(before[1]['Vendor -As per MDM'], 'Test Vendor')
        self.assertNotIn('Source', before[1])
        self.store.reset(str(self.root))
        with report.open(encoding='utf-8-sig') as stream:
            after = list(csv.DictReader(stream))
        self.assertEqual(after, [])
        self.assertEqual(len(report.read_text(encoding='utf-8-sig').splitlines()), 1)

    def test_protected_path_refused(self):
        with patch('services.local_storage.Path.home', return_value=self.root):
            with self.assertRaises(ValueError):
                self.store.reset(str(self.root))
        self.assertEqual(self.store.count(), (1, 1))

    def test_link_preflight_prevents_any_deletion(self):
        # Simulate a reparse point without requiring Windows symlink privileges.
        original = Path.is_junction
        def junction(path):
            return path.name == 'uploads' or original(path)
        with patch.object(Path, 'is_junction', junction):
            with self.assertRaises(ValueError):
                self.store.reset(str(self.root))
        self.assertEqual(self.store.count(), (1, 1))

    def test_locked_file_reports_incomplete_and_retains_database(self):
        original = Path.unlink
        def locked(path, *args, **kwargs):
            if path.name == '0001_test.pdf':
                raise PermissionError('Upload is open')
            return original(path, *args, **kwargs)
        with patch.object(Path, 'unlink', locked):
            with self.assertRaisesRegex(OSError, 'Reset incomplete'):
                self.store.reset(str(self.root))
        self.assertEqual(self.store.count(), (1, 1))



if __name__ == '__main__':
    unittest.main()
