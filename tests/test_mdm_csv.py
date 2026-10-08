import csv
import io
import zipfile
from dataclasses import replace
from datetime import date
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import web_app
from models.invoice_schema import Invoice, InvoiceResult
from services.local_storage import LocalStorage
from services.csv_service import _invoice_age
from services.mdm_csv_service import COLUMNS, DISPLAY_COLUMNS, REPORT_NAME, _received_date, generate_mdm_csv, mdm_row
from test_msg_uploads import msg_bytes


def example(number='000123'):
    return InvoiceResult(invoice=Invoice(
        invoice_number=number, vendor_name='Example Vendor', bank_account_number='001234567',
        invoice_date='2026-10-01', currency='USD', subtotal='100.00', tax_amount='0.00',
        total_amount='100.00', service_start_date='2026-10-01', service_end_date='2026-10-31',
        invoice_comments='Printed invoice note'), processed_at='2026-10-05T00:00:00Z')


def rows(data):
    return list(csv.DictReader(io.StringIO(data.decode('utf-8-sig'))))


def test_mdm_schema_mapping_and_blank_workflow_fields():
    result = rows(generate_mdm_csv([example()]))[0]
    assert list(result) == COLUMNS
    assert COLUMNS == [
        'Business Unit', 'Vendor -As per MDM', 'Billing Account',
        'Invoice Number (As per Invoice)', 'Invoice Number', 'Invoice Date',
        'Currency', 'Billing Start Date', 'Billing End Date', 'Net Amount',
        'Tax Amount', 'Invoice Received Date', 'Sent for Scanning Date',
        'Owner', 'Scanned Date in Expense', 'Ageing (In Days)',
        'Ageing (Today-Sent For Scanning Date)', 'Comments',
    ]
    assert result['Business Unit'] == ''
    assert result['Vendor -As per MDM'] == 'Example Vendor'
    assert result['Invoice Number (As per Invoice)'] == result['Invoice Number'] == '000123'
    assert result['Billing Account'] == ''
    missing_account = example()
    missing_account.invoice.bank_account_number = None
    assert rows(generate_mdm_csv([missing_account]))[0]['Billing Account'] == ''
    assert result['Billing Start Date'] == '10/01/2026'
    assert result['Billing End Date'] == '10/31/2026'
    assert result['Net Amount'] == '100.0'
    assert result['Tax Amount'] == '0.0'
    assert result['Invoice Received Date'] == _received_date('2026-10-05T00:00:00Z')
    assert result['Ageing (In Days)'] == str((date.today() - date(2026, 10, 1)).days)
    assert result['Owner'] == ''
    assert result['Sent for Scanning Date'] == result['Scanned Date in Expense'] == _received_date(example().processed_at)
    assert result['Ageing (Today-Sent For Scanning Date)'] == str(_invoice_age(_received_date(example().processed_at)))
    assert mdm_row(example())['Extraction completeness (%)'] == 100.0
    assert 'Verified accuracy (%)' not in result
    assert result['Comments'] == 'Printed invoice note'
    assert mdm_row(example())['Action'] == 'Review Required'
    assert mdm_row(example())['Review status'] == 'REVIEW_REQUIRED'
    assert 'Review reasons' not in result


def test_review_reasons_name_only_missing_required_fields_and_actual_gates():
    draft = example()
    draft.processing_status = 'LOCAL_DRAFT'
    draft.invoice.bank_account_number = None
    draft.invoice.vendor_email = None
    row = mdm_row(draft)
    assert 'Local extraction has no calibrated field confidence' in row['Review reasons']
    assert 'Billing Account' not in row['Review reasons']
    assert 'Source email' not in row['Review reasons']

    draft.invoice.invoice_date = None
    draft.errors = ['Missing required field: invoice_date']
    row = mdm_row(draft)
    assert 'Missing required: Invoice Date' in row['Review reasons']
    assert 'Local extraction has no calibrated field confidence' in row['Review reasons']

    draft.invoice.invoice_date = '2026-10-01'
    draft.errors = []
    draft.review_required = False
    draft.validation_status = 'VALID'
    assert mdm_row(draft)['Review reasons'] == ''

    draft.duplicate_status = 'POSSIBLE_DUPLICATE'
    draft.review_required = True
    draft.validation_status = 'REVIEW_REQUIRED'
    assert 'Possible duplicate invoice' in mdm_row(draft)['Review reasons']


def test_received_date_is_intake_day_and_ageing_is_today_minus_invoice_date():
    invoice = example()
    invoice.invoice.service_start_date = '2026-09-01'
    invoice.processed_at = '2026-10-05T00:00:00Z'
    first = rows(generate_mdm_csv([invoice]))[0]
    assert first['Invoice Received Date'] == _received_date(invoice.processed_at)
    assert first['Invoice Received Date'] != first['Billing Start Date']
    assert _invoice_age('2026-10-01', date(2026, 10, 7)) == 6
    assert _invoice_age('2026-10-08', date(2026, 10, 7)) == -1
    assert _invoice_age(None, date(2026, 10, 7)) == ''


def test_page_number_is_not_exported_as_billing_account():
    invoice = example()
    invoice.invoice.bank_account_number = 'Page 1'
    row = rows(generate_mdm_csv([invoice]))[0]
    assert row['Billing Account'] == ''
    assert mdm_row(invoice)['Extraction completeness (%)'] == 100.0
    assert row['Comments'] == 'Printed invoice note'


def test_mdm_cumulative_export_preserves_original_csv_and_corrections(tmp_path):
    store = LocalStorage(tmp_path)
    first, second = uuid4().hex, uuid4().hex
    store.save_batch(first, [example('A')])
    store.save_batch(second, [example('B')])
    original = store.root / 'Invoice_Data.csv'
    additional = store.root / REPORT_NAME
    assert original.is_file() and additional.is_file()
    assert 'Date' in rows(original.read_bytes())[0]
    assert 'Business Unit' in rows(additional.read_bytes())[0]
    assert 'Vendor -As per MDM' in rows(additional.read_bytes())[0]
    assert [r['Invoice Number'] for r in rows(additional.read_bytes())] == ['A', 'B']
    assert 'Extraction completeness (%)' not in rows(additional.read_bytes())[0]
    store.save_batch(second, [example('CORRECTED')])
    assert [r['Invoice Number'] for r in rows(additional.read_bytes())] == ['A', 'CORRECTED']
    assert store.count() == (2, 2)
    book = load_workbook(io.BytesIO(store.report_files()['Invoice_History.xlsx']))
    assert [cell.value for cell in book['MDM Invoice Data'][1]] == COLUMNS
    assert book['MDM Invoice Data'].cell(2, COLUMNS.index('Billing Account') + 1).value is None
    assert book['MDM Invoice Data'].cell(2, COLUMNS.index('Net Amount') + 1).value == 100
    book.close()


def test_mdm_batch_history_download_and_zip(tmp_path, monkeypatch):
    monkeypatch.setattr(web_app, 'settings', replace(web_app.settings, local_storage_dir=str(tmp_path)))
    store = LocalStorage(tmp_path)
    batch_id = uuid4().hex
    store.save_batch(batch_id, [example()])
    client = TestClient(web_app.app)
    # A saved report from yesterday must be regenerated before download.
    (store.root / REPORT_NAME).write_bytes(b'stale report')
    for url in [f'/api/download/batch/{batch_id}/{REPORT_NAME}', f'/api/download/history/{REPORT_NAME}']:
        response = client.get(url)
        assert response.status_code == 200
        assert response.headers['content-disposition'] == 'attachment; filename="Invoice_data.csv"'
        assert rows(response.content)[0]['Invoice Number'] == '000123'
    archive = zipfile.ZipFile(io.BytesIO(client.get(f'/api/download/batch/{batch_id}/invoice_results.zip').content))
    assert 'Invoice_Data.csv' in archive.namelist() and REPORT_NAME in archive.namelist()
    assert client.get('/api/download/history/MDM/not-a-report.csv').status_code == 404


def test_empty_mdm_csv_has_exact_headers():
    reader = csv.DictReader(io.StringIO(generate_mdm_csv([]).decode('utf-8-sig')))
    assert reader.fieldnames == COLUMNS
    assert list(reader) == []


def test_failed_scan_has_empty_comments_and_failure_status():
    failed = InvoiceResult.failed('scan.jpg', 'Gemini model temporarily unavailable (HTTP 503)',
                                  '2026-10-05T00:00:00Z')
    row = rows(generate_mdm_csv([failed]))[0]
    assert not row['Invoice Number']
    assert row['Comments'] == ''
    assert mdm_row(failed)['Source'] == 'scan.jpg'
    assert mdm_row(failed)['Action'] == 'Review Required'
    assert mdm_row(failed)['Processing result'] == 'FAILED'
    assert mdm_row(failed)['Review reasons'] == 'Gemini model temporarily unavailable (HTTP 503)'


def test_completed_local_extraction_is_success_but_still_requires_review():
    draft = example()
    draft.processing_status = 'LOCAL_DRAFT'
    draft.validation_status = 'REVIEW_REQUIRED'
    row = mdm_row(draft)
    assert row['Processing result'] == 'SUCCESS'
    assert row['Review status'] == 'REVIEW_REQUIRED'
    assert row['Action'] == 'Review Required'

    ocr_only = example()
    ocr_only.processing_status = 'OCR_ONLY'
    assert mdm_row(ocr_only)['Processing result'] == 'OCR_ONLY'


def test_validated_invoice_has_success_and_no_review_action():
    verified = example()
    verified.validation_status = 'VALID'
    verified.review_required = False
    verified.invoice.vendor_email = 'billing@example.com'
    verified.source_email = 'uploaded.msg'
    row = mdm_row(verified)
    assert row['Review status'] == 'SUCCESS'
    assert row['Action'] == 'No Review Required'
    assert row['Source email'] == 'billing@example.com'
    assert row['Email file'] == 'uploaded.msg'


def test_saved_synthetic_email_is_removed_from_regenerated_report():
    saved = example()
    saved.invoice.vendor_email = '0@example.test'
    assert mdm_row(saved)['Source email'] == ''


@pytest.mark.parametrize('intake', ['upload', 'outlook_folder'])
def test_scanned_attachment_flows_into_cumulative_mdm_csv(tmp_path, monkeypatch, intake):
    from services import document_processing

    root = tmp_path / 'local_data'
    monkeypatch.setattr(web_app, 'settings', replace(web_app.settings,
                        local_storage_dir=str(root)))
    payload = {'invoice': {
        'invoice_number': 'SCAN-001', 'vendor_name': 'Test Vendor',
        'invoice_date': '2026-10-01', 'currency': 'USD', 'subtotal': 100,
        'tax_amount': 5, 'total_amount': 105, 'service_start_date': '2026-10-01',
        'service_end_date': '2026-10-31', 'invoice_comments': 'Printed test invoice',
        'bank_account_number': '000456789',
    }, 'line_items': [], 'field_confidence': {}}
    payload['field_confidence'] = {key: 99 for key in payload['invoice']}
    monkeypatch.setattr(document_processing, 'extract_invoice', lambda data, name: payload)
    monkeypatch.setattr(web_app, 'extraction_ready', lambda: True)
    client = TestClient(web_app.app)
    if intake == 'upload':
        response = client.post('/api/process', files=[('files', ('scan.jpg', b'IMAGE', 'image/jpeg'))])
    else:
        folder = LocalStorage(root).outlook_folder
        folder.joinpath('scan.msg').write_bytes(msg_bytes([{'name': 'scan.jpg', 'data': b'IMAGE'}]))
        response = client.post('/api/outlook/run', json={'tolerance': 0.02, 'threshold': 95})
    assert response.status_code == 200, response.text
    assert response.json()['counts']['total'] == 1
    path = root / REPORT_NAME
    row = rows(path.read_bytes())[0]
    assert row['Invoice Number (As per Invoice)'] == 'SCAN-001'
    assert row['Vendor -As per MDM'] == 'Test Vendor'
    assert row['Billing Account'] == ''
    assert row['Invoice Date'] == '10/01/2026'
    assert row['Currency'] == 'USD'
    assert row['Net Amount'] == '100.0'
    assert row['Tax Amount'] == '5.0'
    assert row['Billing End Date'] == '10/31/2026'
    assert 'Extraction completeness (%)' not in row
    assert response.json()['results'][0]['mdm_row']['Source'] in {'scan.jpg', 'scan.msg::scan.jpg'}
    assert 'Printed test invoice' in row['Comments']
    assert rows(client.get(f'/api/download/history/{REPORT_NAME}').content)[0] == row
