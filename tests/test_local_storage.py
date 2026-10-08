"""Persistence regression tests using isolated local temporary directories."""
import csv
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from uuid import uuid4

from openpyxl import load_workbook

from models.invoice_schema import Invoice, InvoiceResult, LineItem
from services.local_storage import LocalStorage, _atomic_write
from utils.file_utils import DocumentInput


def result(number='INV-01'):
    return InvoiceResult(invoice=Invoice(invoice_number=number, vendor_name='Example',
                                        bank_account_number='0012345', total_amount='123.40'),
                         line_items=[LineItem(description='Service', quantity=1, line_total='123.40')],
                         source_file='invoice.pdf', processed_at='2026-10-05T12:00:00+00:00')


class LocalStorageTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.store = LocalStorage(self.root)

    def rows(self):
        return list(csv.DictReader(io.StringIO((self.root / 'Invoice_Data.csv').read_text(encoding='utf-8-sig'))))

    def test_append_and_reopen_all_reports(self):
        self.assertEqual(self.store.save_batch(uuid4().hex, [result('FIRST')]), [])
        self.assertEqual(LocalStorage(self.root).save_batch(uuid4().hex, [result('SECOND')]), [])
        self.assertEqual([row['Invoice'] for row in self.rows()], ['FIRST', 'SECOND'])
        self.assertEqual(self.store.count(), (2, 2))
        book = load_workbook(io.BytesIO(self.store.report_files()['Invoice_History.xlsx']))
        self.assertEqual(book.sheetnames, ['Invoice_Data', 'Invoice_Line_Items', 'Audit_Report',
                                           'MDM Invoice Data'])
        for sheet in book:
            self.assertEqual(sheet.max_row, 3)
            self.assertEqual(sheet.freeze_panes, 'A2')
        headers = [cell.value for cell in book['Invoice_Data'][1]]
        self.assertEqual(book['Invoice_Data'].cell(2, headers.index('MDM Account number') + 1).value, '0012345')
        self.assertEqual(book['Invoice_Data'].cell(2, headers.index('Invoice Amount') + 1).value, 123.4)
        mdm_headers = [cell.value for cell in book['MDM Invoice Data'][1]]
        self.assertIsNone(book['MDM Invoice Data'].cell(2, mdm_headers.index('Billing Account') + 1).value)
        book.close()

    def test_same_batch_retry_and_correction_dont_append(self):
        batch = uuid4().hex
        self.store.save_batch(batch, [result()])
        self.store.save_batch(batch, [result()])
        corrected = result('CORRECTED')
        self.store.save_batch(batch, [corrected])
        self.assertEqual(self.store.count(), (1, 1))
        self.assertEqual([row['Invoice'] for row in self.rows()], ['CORRECTED'])

    def test_originals_and_documents_do_not_overwrite_or_escape(self):
        first, second = uuid4().hex, uuid4().hex
        self.store.archive_uploads(first, [('../../invoice.pdf', b'A'), ('invoice.pdf', b'B')])
        self.store.archive_uploads(second, [('invoice.pdf', b'C')])
        self.store.archive_documents(first, [DocumentInput('invoice.pdf', b'D', source_email='mail.msg')])
        files = sorted((self.root / 'uploads' / first / 'originals').glob('*.pdf'))
        self.assertEqual([p.read_bytes() for p in files], [b'A', b'B'])
        self.assertEqual(len(list((self.root / 'uploads').rglob('*.pdf'))), 4)
        self.assertFalse((self.root / 'invoice.pdf').exists())
        with self.assertRaises(ValueError):
            self.store.archive_uploads('../escape', [])

    def test_locked_excel_preserves_history_and_retries(self):
        self.store.save_batch(uuid4().hex, [result('OLD')])
        previous = (self.root / 'Invoice_History.xlsx').read_bytes()

        def locked(path, payload):
            if path.suffix == '.xlsx':
                raise PermissionError('Workbook is open in Excel')
            _atomic_write(path, payload)

        batch = uuid4().hex
        with patch('services.local_storage._atomic_write', side_effect=locked):
            warnings = self.store.save_batch(batch, [result('NEW')])
        self.assertTrue(warnings)
        self.assertTrue(LocalStorage(self.root).export_warnings())
        self.assertEqual(self.store.count(), (2, 2))
        self.assertEqual((self.root / 'Invoice_History.xlsx').read_bytes(), previous)
        self.assertEqual(self.store.refresh_exports(), [])
        self.assertEqual(self.store.export_warnings(), [])
        self.assertEqual(len(self.rows()), 2)

    def test_failed_and_skipped_records_and_msg_provenance(self):
        failed = InvoiceResult.failed('broken.pdf', 'test error', '2026-10-05T12:00:00+00:00')
        failed.source_email = 'mail.msg'
        failed.source_attachment = 'broken.pdf'
        failed.attachment_index = 2
        self.store.save_batch(uuid4().hex, [failed])
        self.assertEqual(self.rows()[0]['source_email'], 'mail.msg')
        self.assertEqual(self.rows()[0]['Status'], 'FAILED')

    def test_excel_text_is_not_executable_formula(self):
        entry = result('=1+2')
        self.store.save_batch(uuid4().hex, [entry])
        book = load_workbook(io.BytesIO(self.store.report_files()['Invoice_History.xlsx']))
        cell = book['Invoice_Data']['B2']
        self.assertEqual(cell.value, '=1+2')
        self.assertEqual(cell.data_type, 's')
        book.close()

    def test_empty_history_export(self):
        self.store.refresh_exports()
        self.assertEqual(self.rows(), [])
        self.assertEqual(len(self.store.report_files()), 5)



if __name__ == '__main__':
    unittest.main()
